import json
from dataclasses import dataclass
from io import BytesIO

import openpyxl
from openpyxl.styles import Font


class WorkbookBuildError(Exception):
    pass


@dataclass
class CellSpec:
    value: str | int | float | None = None
    formula: str | None = None
    bold: bool = False
    number_format: str | None = None


@dataclass
class SheetSpec:
    name: str
    columns: list[str]
    rows: list[list]
    column_widths: dict[str, int] | None = None


@dataclass
class WorkbookSpec:
    sheets: list[SheetSpec]
    title: str | None = None


def parse_workbook_spec(json_str: str) -> WorkbookSpec:
    try:
        data = json.loads(json_str)
    except json.JSONDecodeError as e:
        raise WorkbookBuildError(f"Invalid JSON: {e}") from e

    if not isinstance(data, dict) or "sheets" not in data:
        raise WorkbookBuildError("JSON must contain a 'sheets' key")

    sheets_data = data["sheets"]
    if not sheets_data:
        raise WorkbookBuildError("Workbook must have at least one sheet")

    sheets = []
    for sheet_data in sheets_data:
        if "columns" not in sheet_data:
            raise WorkbookBuildError(
                f"Sheet '{sheet_data.get('name', '?')}' missing required 'columns' field"
            )

        rows = []
        for raw_row in sheet_data.get("rows", []):
            row = []
            for cell in raw_row:
                if isinstance(cell, dict):
                    row.append(CellSpec(
                        value=cell.get("value"),
                        formula=cell.get("formula"),
                        bold=cell.get("bold", False),
                        number_format=cell.get("number_format"),
                    ))
                else:
                    row.append(cell)
            rows.append(row)

        column_widths = sheet_data.get("column_widths")
        sheets.append(SheetSpec(
            name=sheet_data.get("name", "Sheet1"),
            columns=sheet_data["columns"],
            rows=rows,
            column_widths=column_widths,
        ))

    return WorkbookSpec(sheets=sheets, title=data.get("title"))


def build_workbook_from_spec(spec: WorkbookSpec) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_spec in spec.sheets:
        ws = wb.create_sheet(title=sheet_spec.name)
        _write_sheet(ws, sheet_spec)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def modify_workbook(original_bytes: bytes, spec: WorkbookSpec) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(original_bytes))

    # Remove sheets that will be replaced
    for sheet_spec in spec.sheets:
        if sheet_spec.name in wb.sheetnames:
            del wb[sheet_spec.name]

    # Add new/replaced sheets
    for sheet_spec in spec.sheets:
        ws = wb.create_sheet(title=sheet_spec.name)
        _write_sheet(ws, sheet_spec)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_sheet(ws, sheet_spec: SheetSpec) -> None:
    # Write headers with bold
    bold_font = Font(bold=True)
    for col_idx, header in enumerate(sheet_spec.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font

    # Write data rows
    for row_idx, row_data in enumerate(sheet_spec.rows, 2):
        for col_idx, cell_data in enumerate(row_data, 1):
            if isinstance(cell_data, CellSpec):
                if cell_data.formula:
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_data.formula)
                else:
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_data.value)
                if cell_data.bold:
                    cell.font = Font(bold=True)
                if cell_data.number_format:
                    cell.number_format = cell_data.number_format
            else:
                ws.cell(row=row_idx, column=col_idx, value=cell_data)

    # Set column widths
    if sheet_spec.column_widths:
        for col_letter, width in sheet_spec.column_widths.items():
            ws.column_dimensions[col_letter].width = width
