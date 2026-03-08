from dataclasses import dataclass, field
from io import BytesIO

import openpyxl


class ExcelReadError(Exception):
    pass


@dataclass
class SheetSummary:
    name: str
    headers: list[str]
    rows: list[list[str]]
    total_rows: int
    truncated: bool


@dataclass
class SpreadsheetSummary:
    filename: str
    sheets: list[SheetSummary] = field(default_factory=list)


def read_spreadsheet(
    file_bytes: bytes, filename: str, max_rows: int = 50, max_size_mb: int = 10
) -> SpreadsheetSummary:
    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > max_size_mb:
        raise ExcelReadError(
            f"File '{filename}' is {size_mb:.1f} MB, which exceeds the {max_size_mb} MB limit."
        )

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ExcelReadError(f"Could not read '{filename}': {e}") from e

    summary = SpreadsheetSummary(filename=filename)

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows = []
            for row in ws.iter_rows(values_only=True):
                all_rows.append(row)

            if not all_rows:
                summary.sheets.append(
                    SheetSummary(
                        name=sheet_name,
                        headers=[],
                        rows=[],
                        total_rows=0,
                        truncated=False,
                    )
                )
                continue

            headers = [_cell_to_str(c) for c in all_rows[0]]
            data_rows = all_rows[1:]
            total_rows = len(data_rows)
            truncated = total_rows > max_rows
            preview_rows = data_rows[:max_rows]

            summary.sheets.append(
                SheetSummary(
                    name=sheet_name,
                    headers=headers,
                    rows=[[_cell_to_str(c) for c in row] for row in preview_rows],
                    total_rows=total_rows,
                    truncated=truncated,
                )
            )
    finally:
        wb.close()

    return summary


def format_summary_for_prompt(summary: SpreadsheetSummary) -> str:
    parts = [f"## Spreadsheet: {summary.filename}"]

    for sheet in summary.sheets:
        parts.append(f"\n### Sheet: {sheet.name}")

        if not sheet.headers and not sheet.rows:
            parts.append("(empty sheet)")
            continue

        if sheet.headers:
            parts.append("| " + " | ".join(sheet.headers) + " |")
            parts.append("| " + " | ".join("---" for _ in sheet.headers) + " |")

        for row in sheet.rows:
            if len(row) < len(sheet.headers):
                padded = row + [""] * (len(sheet.headers) - len(row))
            else:
                padded = row
            parts.append("| " + " | ".join(padded[: len(sheet.headers)]) + " |")

        if sheet.truncated:
            parts.append(
                f"\n*Showing {len(sheet.rows)} of {sheet.total_rows} data rows (truncated)*"
            )

    return "\n".join(parts)


MAX_CELL_LENGTH = 1000


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if len(text) > MAX_CELL_LENGTH:
        text = text[:MAX_CELL_LENGTH] + "..."
    # Escape pipes to prevent breaking markdown table formatting
    text = text.replace("|", "\\|")
    return text
