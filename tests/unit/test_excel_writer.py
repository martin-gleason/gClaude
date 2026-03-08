import json
from io import BytesIO

import openpyxl
import pytest

from src.excel.writer import (
    CellSpec,
    SheetSpec,
    WorkbookBuildError,
    WorkbookSpec,
    build_workbook_from_spec,
    modify_workbook,
    parse_workbook_spec,
)


def _make_simple_spec_json(**overrides):
    spec = {
        "sheets": [
            {
                "name": "Sheet1",
                "columns": ["Name", "Amount"],
                "rows": [["Alice", 100], ["Bob", 200]],
            }
        ],
    }
    spec.update(overrides)
    return json.dumps(spec)


def _read_xlsx(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(data))


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Helper: create a real xlsx with given data for each sheet."""
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseWorkbookSpec:
    def test_parse_simple_spec(self):
        spec = parse_workbook_spec(_make_simple_spec_json())
        assert len(spec.sheets) == 1
        assert spec.sheets[0].name == "Sheet1"
        assert spec.sheets[0].columns == ["Name", "Amount"]
        assert len(spec.sheets[0].rows) == 2

    def test_parse_with_formulas(self):
        json_str = json.dumps({
            "sheets": [{
                "name": "Sheet1",
                "columns": ["Item", "Total"],
                "rows": [
                    ["Widgets", {"formula": "=SUM(B2:B10)"}],
                ],
            }]
        })
        spec = parse_workbook_spec(json_str)
        cell = spec.sheets[0].rows[0][1]
        assert isinstance(cell, CellSpec)
        assert cell.formula == "=SUM(B2:B10)"

    def test_parse_with_formatting(self):
        json_str = json.dumps({
            "sheets": [{
                "name": "Sheet1",
                "columns": ["Item", "Price"],
                "rows": [
                    ["Widgets", {"value": 9.99, "bold": True, "number_format": "#,##0.00"}],
                ],
            }]
        })
        spec = parse_workbook_spec(json_str)
        cell = spec.sheets[0].rows[0][1]
        assert isinstance(cell, CellSpec)
        assert cell.bold is True
        assert cell.number_format == "#,##0.00"
        assert cell.value == 9.99

    def test_parse_multi_sheet(self):
        json_str = json.dumps({
            "sheets": [
                {"name": "Income", "columns": ["Source"], "rows": [["Salary"]]},
                {"name": "Expenses", "columns": ["Category"], "rows": [["Rent"]]},
            ]
        })
        spec = parse_workbook_spec(json_str)
        assert len(spec.sheets) == 2
        assert spec.sheets[0].name == "Income"
        assert spec.sheets[1].name == "Expenses"

    def test_parse_invalid_json_raises(self):
        with pytest.raises(WorkbookBuildError, match="Invalid JSON"):
            parse_workbook_spec("not valid json {{{")

    def test_parse_missing_columns_raises(self):
        json_str = json.dumps({
            "sheets": [{"name": "Sheet1", "rows": [["a"]]}]
        })
        with pytest.raises(WorkbookBuildError, match="columns"):
            parse_workbook_spec(json_str)

    def test_parse_empty_sheets_raises(self):
        json_str = json.dumps({"sheets": []})
        with pytest.raises(WorkbookBuildError, match="at least one sheet"):
            parse_workbook_spec(json_str)


class TestBuildWorkbookFromSpec:
    def test_build_single_sheet(self):
        spec = parse_workbook_spec(_make_simple_spec_json())
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        assert len(wb.sheetnames) == 1
        assert wb.sheetnames[0] == "Sheet1"

    def test_build_with_headers(self):
        spec = parse_workbook_spec(_make_simple_spec_json())
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        ws = wb["Sheet1"]
        assert ws.cell(1, 1).value == "Name"
        assert ws.cell(1, 2).value == "Amount"

    def test_build_with_formulas(self):
        json_str = json.dumps({
            "sheets": [{
                "name": "Sheet1",
                "columns": ["Item", "Total"],
                "rows": [
                    ["Widgets", {"formula": "=SUM(B2:B10)"}],
                ],
            }]
        })
        spec = parse_workbook_spec(json_str)
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        ws = wb["Sheet1"]
        assert ws.cell(2, 2).value == "=SUM(B2:B10)"

    def test_build_with_bold_header(self):
        spec = parse_workbook_spec(_make_simple_spec_json())
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        ws = wb["Sheet1"]
        assert ws.cell(1, 1).font.bold is True

    def test_build_with_number_format(self):
        json_str = json.dumps({
            "sheets": [{
                "name": "Sheet1",
                "columns": ["Price"],
                "rows": [
                    [{"value": 9.99, "number_format": "#,##0.00"}],
                ],
            }]
        })
        spec = parse_workbook_spec(json_str)
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        ws = wb["Sheet1"]
        assert ws.cell(2, 1).number_format == "#,##0.00"

    def test_build_multi_sheet(self):
        json_str = json.dumps({
            "sheets": [
                {"name": "Income", "columns": ["Source"], "rows": [["Salary"]]},
                {"name": "Expenses", "columns": ["Item"], "rows": [["Rent"]]},
            ]
        })
        spec = parse_workbook_spec(json_str)
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        assert wb.sheetnames == ["Income", "Expenses"]

    def test_build_column_widths(self):
        json_str = json.dumps({
            "sheets": [{
                "name": "Sheet1",
                "columns": ["Name", "Amount"],
                "rows": [["Alice", 100]],
                "column_widths": {"A": 25, "B": 15},
            }]
        })
        spec = parse_workbook_spec(json_str)
        data = build_workbook_from_spec(spec)
        wb = _read_xlsx(data)
        ws = wb["Sheet1"]
        assert ws.column_dimensions["A"].width == 25
        assert ws.column_dimensions["B"].width == 15

    def test_roundtrip(self):
        """Build xlsx, then read back with reader.py and verify data matches."""
        from src.excel.reader import read_spreadsheet

        spec = parse_workbook_spec(_make_simple_spec_json())
        data = build_workbook_from_spec(spec)
        summary = read_spreadsheet(data, "test.xlsx")
        assert len(summary.sheets) == 1
        sheet = summary.sheets[0]
        assert sheet.headers == ["Name", "Amount"]
        assert sheet.rows[0][0] == "Alice"
        assert sheet.rows[0][1] == "100"
        assert sheet.rows[1][0] == "Bob"
        assert sheet.rows[1][1] == "200"


class TestModifyWorkbook:
    def test_add_new_sheet_to_existing(self):
        original = _make_xlsx({"Data": [["A", "B"], [1, 2]]})
        spec = WorkbookSpec(sheets=[
            SheetSpec(name="Summary", columns=["Total"], rows=[[42]]),
        ])
        result = modify_workbook(original, spec)
        wb = _read_xlsx(result)
        assert "Data" in wb.sheetnames
        assert "Summary" in wb.sheetnames
        # Original data preserved
        assert wb["Data"].cell(1, 1).value == "A"

    def test_replace_existing_sheet(self):
        original = _make_xlsx({"Data": [["Old", "Values"], [1, 2]]})
        spec = WorkbookSpec(sheets=[
            SheetSpec(name="Data", columns=["New", "Columns"], rows=[[10, 20]]),
        ])
        result = modify_workbook(original, spec)
        wb = _read_xlsx(result)
        assert wb["Data"].cell(1, 1).value == "New"
        assert wb["Data"].cell(2, 1).value == 10

    def test_original_unmodified_sheets_preserved(self):
        original = _make_xlsx({
            "Keep": [["X"], [1]],
            "Replace": [["Y"], [2]],
        })
        spec = WorkbookSpec(sheets=[
            SheetSpec(name="Replace", columns=["Z"], rows=[[3]]),
        ])
        result = modify_workbook(original, spec)
        wb = _read_xlsx(result)
        assert wb["Keep"].cell(1, 1).value == "X"
        assert wb["Keep"].cell(2, 1).value == 1
        assert wb["Replace"].cell(1, 1).value == "Z"
