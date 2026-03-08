from io import BytesIO

import openpyxl
import pytest

from src.excel.analyzer import (
    AnalysisError,
    analyze_spreadsheet,
    format_analysis_for_prompt,
)


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Helper to create an in-memory xlsx file from sheet data."""
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestAnalyzeSpreadsheet:
    def test_basic_analysis(self):
        data = _make_xlsx({"Sheet1": [["Name", "Age"], ["Alice", 30], ["Bob", 25]]})
        result = analyze_spreadsheet(data, "test.xlsx")
        assert len(result.sheets) == 1
        sheet = result.sheets[0]
        assert sheet.name == "Sheet1"
        assert sheet.row_count == 2
        assert sheet.column_count == 2
        assert "Name" in sheet.columns
        assert "Age" in sheet.columns

    def test_numeric_summary_computed(self):
        data = _make_xlsx({"Sheet1": [["Value"], [10], [20], [30]]})
        result = analyze_spreadsheet(data, "test.xlsx")
        sheet = result.sheets[0]
        assert "Value" in sheet.numeric_summary
        stats = sheet.numeric_summary["Value"]
        assert "mean" in stats
        assert stats["mean"] == 20.0

    def test_multi_sheet_analysis(self):
        data = _make_xlsx({
            "Sales": [["Amount"], [100], [200]],
            "Costs": [["Amount"], [50], [75]],
        })
        result = analyze_spreadsheet(data, "test.xlsx")
        assert len(result.sheets) == 2
        names = [s.name for s in result.sheets]
        assert "Sales" in names
        assert "Costs" in names

    def test_dtypes_captured(self):
        data = _make_xlsx({"Sheet1": [["Name", "Age"], ["Alice", 30]]})
        result = analyze_spreadsheet(data, "test.xlsx")
        sheet = result.sheets[0]
        assert "Name" in sheet.dtypes
        assert "Age" in sheet.dtypes

    def test_sample_values_included(self):
        data = _make_xlsx({"Sheet1": [["Name"], ["Alice"], ["Bob"], ["Charlie"]]})
        result = analyze_spreadsheet(data, "test.xlsx")
        sheet = result.sheets[0]
        assert "Name" in sheet.sample_values
        assert "Alice" in sheet.sample_values["Name"]

    def test_empty_sheet(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Empty"
        buf = BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        result = analyze_spreadsheet(data, "test.xlsx")
        assert len(result.sheets) == 1
        assert result.sheets[0].row_count == 0

    def test_file_too_large_raises_error(self):
        with pytest.raises(AnalysisError, match="exceeds"):
            analyze_spreadsheet(b"x" * (11 * 1024 * 1024), "big.xlsx", max_size_mb=10)

    def test_corrupt_file_raises_error(self):
        with pytest.raises(AnalysisError, match="Could not analyze"):
            analyze_spreadsheet(b"not-an-xlsx-file", "bad.xlsx")

    def test_max_rows_honored(self):
        rows = [["Value"]] + [[i] for i in range(100)]
        data = _make_xlsx({"Sheet1": rows})
        result = analyze_spreadsheet(data, "test.xlsx", max_rows=10)
        assert result.sheets[0].row_count == 10


class TestFormatAnalysisForPrompt:
    def test_format_includes_filename(self):
        data = _make_xlsx({"Sheet1": [["A"], [1]]})
        analysis = analyze_spreadsheet(data, "report.xlsx")
        text = format_analysis_for_prompt(analysis)
        assert "report.xlsx" in text

    def test_format_includes_statistics(self):
        data = _make_xlsx({"Sheet1": [["Value"], [10], [20], [30]]})
        analysis = analyze_spreadsheet(data, "test.xlsx")
        text = format_analysis_for_prompt(analysis)
        assert "mean" in text
        assert "20.00" in text
